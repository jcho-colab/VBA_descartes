import openpyxl
import pandas as pd
from dataclasses import dataclass
from typing import Dict, List, Any, Optional
import logging
import os

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

@dataclass
class AppConfig:
    country: str
    year: str
    min_chapter: int
    max_csv: int
    zd14_date: Any
    rate_type_defs: pd.DataFrame
    uom_dict: Dict[str, str]
    country_list: List[str]
    chapter_list: List[str]
    active_country_group_list: List[str]
    all_country_group_list: List[str]

class ConfigLoader:
    def __init__(self, excel_path: str):
        self.excel_path = excel_path
        self.wb = None
        
    def load(self, country_override: str = None) -> AppConfig:
        logger.info(f"Loading configuration from {self.excel_path}")
        
        # Check if file exists
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Configuration file not found: {self.excel_path}")
        
        self.wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        
        # Load Menu settings
        menu_sheet = self.wb["Menu"]
        
        if country_override:
            country = country_override
            logger.info(f"Using Country Override: {country}")
        else:
            country = self._get_named_range_value("Country")
            if not country:
                raise ValueError("Country not found in configuration")
            
        year = str(self._get_named_range_value("Year")) 
        min_chapter = int(self._get_named_range_value("MinChapter") or 25)
        max_csv = int(self._get_named_range_value("MaxCSV") or 1000000)
        zd14_date = self._get_named_range_value("ZD14Date")
        
        logger.info(f"Loaded Global Settings: Country={country}, Year={year}, MinChapter={min_chapter}")
        
        # Generate chapter list (25-99 typically)
        chapter_list = [str(i).zfill(2) for i in range(min_chapter, 100)]
        
        # Load Config tables
        config_sheet = self.wb["Config"]
        
        # Helper to find table case-insensitively
        def find_table(base_name):
            # Try exact
            if base_name in config_sheet.tables:
                return base_name
            # Try case-insensitive match
            for key in config_sheet.tables.keys():
                if key.lower() == base_name.lower():
                    return key
            return None

        # Load RateType Table
        rate_type_base = f"{country}RateType"
        rate_type_table_name = find_table(rate_type_base)
        
        if not rate_type_table_name:
            logger.warning(f"RateType table '{rate_type_base}' not found")
            rate_type_df = pd.DataFrame()
        else:
            rate_type_df = self._read_table(config_sheet, rate_type_table_name)
        
        # Load UOM Table
        uom_base = f"{country}UOM"
        uom_table_name = find_table(uom_base)
        
        if not uom_table_name:
            logger.warning(f"UOM table '{uom_base}' not found")
            uom_df = pd.DataFrame()
        else:
            uom_df = self._read_table(config_sheet, uom_table_name)
        
        uom_dict = {}
        if not uom_df.empty and "Descartes UOM" in uom_df.columns and "SAP UOM" in uom_df.columns:
             uom_dict = dict(zip(uom_df["Descartes UOM"], uom_df["SAP UOM"]))

        # Load Country List (for EU)
        country_list = [country]
        if country == "EU":
            cl_base = f"{country}CountryList"
            country_list_table_name = find_table(cl_base)
            if country_list_table_name:
                cl_df = self._read_table(config_sheet, country_list_table_name)
                if not cl_df.empty:
                    country_list = cl_df.iloc[:, 0].tolist()
        
        # Extract active and all country group lists
        active_country_group_list = []
        all_country_group_list = []
        
        if not rate_type_df.empty and "Descartes CG" in rate_type_df.columns:
            for _, row in rate_type_df.iterrows():
                cg_full = row["Descartes CG"]
                if pd.notna(cg_full):
                    # The "Descartes CG" contains both country_group and duty_rate_type
                    # e.g., "_DNZ1 B001" where "_DNZ1" is country_group and "B001" is duty_rate_type
                    # Split and extract just the country_group part
                    cg_parts = str(cg_full).split()
                    cg = cg_parts[0] if cg_parts else str(cg_full)
                    
                    # Add full version for filtering logic
                    all_country_group_list.append(str(cg_full))
                    
                    # Also add country_group alone for validation
                    if cg not in all_country_group_list:
                        all_country_group_list.append(cg)
                    
                    comment = str(row.get("Comment", "")).lower()
                    if "remove" not in comment:
                        active_country_group_list.append(str(cg_full))
                        if cg not in active_country_group_list:
                            active_country_group_list.append(cg)

        return AppConfig(
            country=country,
            year=year,
            min_chapter=min_chapter,
            max_csv=max_csv,
            zd14_date=zd14_date,
            rate_type_defs=rate_type_df,
            uom_dict=uom_dict,
            country_list=country_list,
            chapter_list=chapter_list,
            active_country_group_list=active_country_group_list,
            all_country_group_list=all_country_group_list
        )

    def _get_named_range_value(self, name: str) -> Any:
        """Retrieves the value of a named range."""
        try:
            defined_name = self.wb.defined_names[name]
            # This returns a list of destinations, usually just one
            for title, coord in defined_name.destinations:
                ws = self.wb[title]
                return ws[coord].value
            return None
        except Exception as e:
            logger.warning(f"Could not read named range '{name}': {e}")
            return None

    def _read_table(self, sheet, table_name: str) -> pd.DataFrame:
        """Reads an Excel table (ListObject) into a DataFrame."""
        try:
            # openpyxl stores tables in sheet.tables
            if table_name not in sheet.tables:
                logger.warning(f"Table '{table_name}' not found in sheet '{sheet.title}'")
                return pd.DataFrame()

            tbl = sheet.tables[table_name]
            ref = tbl.ref # e.g. "A1:C5"
            
            data = sheet[ref]
            rows = [[cell.value for cell in row] for row in data]
            
            if not rows:
                return pd.DataFrame()

            header = rows[0]
            body = rows[1:]
            
            df = pd.DataFrame(body, columns=header)
            # Remove rows where all values are None
            df = df.dropna(how='all')
            
            return df
        except Exception as e:
            logger.error(f"Error reading table '{table_name}': {e}")
            return pd.DataFrame()
